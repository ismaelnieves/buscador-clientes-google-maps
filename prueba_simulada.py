#!/usr/bin/env python3
"""Prueba buscar_clientes.py con datos inventados, sin llamar a Google ni a ninguna web.

Uso:  python prueba_simulada.py
Genera un Excel de muestra en esta carpeta y revisa que los resultados sean los esperados.
"""

import tempfile
from pathlib import Path

import requests
from openpyxl import load_workbook

import buscar_clientes as bc

PAGINA_1 = {
    "places": [
        {"id": "a1", "displayName": {"text": "Carpintería Sin Web"}, "formattedAddress": "Av. Italia 1234, Montevideo",
         "nationalPhoneNumber": "099 123 456", "internationalPhoneNumber": "+598 99 123 456",
         "googleMapsUri": "https://maps.google.com/?cid=1", "rating": 4.7, "userRatingCount": 35,
         "businessStatus": "OPERATIONAL"},
        {"id": "a2", "displayName": {"text": "Muebles Insta"}, "formattedAddress": "Cordón, Montevideo",
         "nationalPhoneNumber": "2400 0000", "internationalPhoneNumber": "+598 2400 0000",
         "websiteUri": "https://www.instagram.com/mueblesinsta", "googleMapsUri": "https://maps.google.com/?cid=2",
         "rating": 4.2, "userRatingCount": 8, "businessStatus": "OPERATIONAL"},
        {"id": "a3", "displayName": {"text": "Carpintería Cerrada"}, "businessStatus": "CLOSED_PERMANENTLY"},
    ],
    "nextPageToken": "TOKEN2",
}
PAGINA_2 = {
    "places": [
        {"id": "a1", "displayName": {"text": "Carpintería Sin Web"}},  # repetido
        {"id": "a4", "displayName": {"text": "Maderas Viejas"}, "nationalPhoneNumber": "098 765 432",
         "internationalPhoneNumber": "+598 98 765 432", "websiteUri": "http://maderasviejas.com.uy",
         "userRatingCount": 50, "businessStatus": "OPERATIONAL"},
        {"id": "a5", "displayName": {"text": "Web Caída SA"}, "nationalPhoneNumber": "2600 1111",
         "internationalPhoneNumber": "+598 2600 1111", "websiteUri": "https://caida.com.uy",
         "businessStatus": "OPERATIONAL"},
        {"id": "a6", "displayName": {"text": "Moderna Ltda"}, "nationalPhoneNumber": "2600 2222",
         "internationalPhoneNumber": "+598 2600 2222", "websiteUri": "https://moderna.com.uy",
         "businessStatus": "OPERATIONAL"},
    ],
}

WEB_VIEJA = ("<html><head><title>Maderas</title></head><body><table width='800'><tr><td>"
             "<font face='Arial'>Bienvenidos</font> Escribinos a ventas@maderasviejas.com.uy"
             "<script src='jquery-1.4.2.min.js'></script></td></tr></table>© 2015 Maderas Viejas</body></html>")
WEB_MODERNA = ("<html><head><meta name='viewport' content='width=device-width'>"
               "<meta name='description' content='Muebles a medida'></head><body>"
               "<a href='https://wa.me/59899000000'>WhatsApp</a><form></form>© 2026 Moderna"
               + " " * 200 + "</body></html>")


class Respuesta:
    def __init__(self, status_code, cuerpo=None, text="", url=""):
        self.status_code, self._cuerpo, self.text, self.url = status_code, cuerpo, text, url

    def json(self):
        return self._cuerpo


class SesionFalsa:
    def __init__(self):
        self.pedidos_places = []

    def post(self, url, headers, json, timeout):
        assert url == bc.URL_PLACES
        assert headers["X-Goog-Api-Key"] == "CLAVE_DE_PRUEBA"
        self.pedidos_places.append(dict(json))
        return Respuesta(200, PAGINA_2 if json.get("pageToken") == "TOKEN2" else PAGINA_1)

    def get(self, url, **kwargs):
        if "caida" in url:
            raise requests.ConnectionError("no responde")
        if "maderasviejas" in url:
            return Respuesta(200, text=WEB_VIEJA, url=url)
        return Respuesta(200, text=WEB_MODERNA, url=url)


def main():
    bc.time.sleep = lambda s: None  # no esperar entre páginas
    config = {**bc.CONFIG_POR_DEFECTO, "google_api_key": "CLAVE_DE_PRUEBA"}
    sesion = SesionFalsa()
    with tempfile.TemporaryDirectory() as tmp:
        ruta = bc.ejecutar("carpinterías", "Montevideo", "Uruguay", config, sesion, tmp)
        wb = load_workbook(ruta)
        ws = wb["Posibles clientes"]
        filas = [dict(zip([c.value for c in ws[1]], [c.value for c in f])) for f in ws.iter_rows(min_row=2)]
        muestra = Path(__file__).resolve().parent / "ejemplo-salida-simulada.xlsx"
        wb.save(muestra)

    assert len(sesion.pedidos_places) == 2, "debía pedir 2 páginas"
    assert sesion.pedidos_places[0]["textQuery"] == "carpinterías en Montevideo, Uruguay"
    assert sesion.pedidos_places[0]["regionCode"] == "UY"
    assert bc.datos_pais("Perú") == bc.datos_pais(" peru ") == ("PE", ("519",)), "acepta con o sin tilde"
    assert bc.datos_pais("Narnia") == (None, ()), "un país desconocido no rompe nada"
    assert bc.es_celular("+54 9 11 5555-1234", bc.datos_pais("Argentina")[1])
    assert not bc.es_celular("+54 11 4444-1234", bc.datos_pais("Argentina")[1])

    otra = SesionFalsa()
    bc.buscar_en_places("pizzerías", "Ciudad Gótica", "Narnia", config, otra)
    assert "regionCode" not in otra.pedidos_places[0], "sin región si no se conoce el país"
    por_nombre = {f["Nombre"]: f for f in filas}
    assert "Carpintería Cerrada" not in por_nombre, "los cerrados se descartan"
    assert "Moderna Ltda" not in por_nombre, "una web correcta queda afuera por puntaje"
    assert sorted(por_nombre) == ["Carpintería Sin Web", "Maderas Viejas", "Muebles Insta", "Web Caída SA"], list(por_nombre)
    assert [f["Puntaje"] for f in filas] == [7, 7, 6, 5], "ordenados de mayor a menor puntaje"

    a = por_nombre["Carpintería Sin Web"]
    assert (a["Estado web"], a["Puntaje"], a["Prioridad"], a["WhatsApp"]) == ("No tiene", 7, "Alta", "099 123 456")
    v = por_nombre["Maderas Viejas"]
    assert v["Estado web"] == "Vieja" and v["Correo"] == "ventas@maderasviejas.com.uy" and v["Puntaje"] == 7
    assert v["Problema principal"] == "Sin celular"
    i = por_nombre["Muebles Insta"]
    assert (i["Estado web"], i["Puntaje"], i["WhatsApp"]) == ("Solo redes o app de terceros", 6, "No figura")
    assert "instagram.com" in i["Plantilla WhatsApp"]
    c = por_nombre["Web Caída SA"]
    assert (c["Estado web"], c["Puntaje"], c["Prioridad"]) == ("Caída o no responde", 5, "Media")

    for f in filas:
        print(f"  {f['Puntaje']:>2}  {f['Prioridad']:<6} {f['Nombre']:<22} {f['Estado web']}")
    print(f"\nTodo bien. Excel de muestra: {muestra}")


if __name__ == "__main__":
    main()
