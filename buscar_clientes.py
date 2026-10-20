#!/usr/bin/env python3
"""Busca posibles clientes en Google Places y arma un Excel con plantillas de mensajes.

Busca negocios de un rubro en una ciudad, revisa la web de cada uno, les pone
un puntaje y escribe plantillas de correo y WhatsApp sin IA.

NO envía nada: el resultado es un Excel para revisar y mandar a mano.

Uso:
    python buscar_clientes.py "carpinterías" "Córdoba" "Argentina"

Si se abre sin datos (por ejemplo con doble clic en buscar_clientes.exe),
pregunta el rubro, la ciudad y el país.
"""

import argparse
import datetime
import json
import re
import sys
import time
import unicodedata
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

# Carpeta del programa: la del .exe si está compilado, si no la del script.
CARPETA = Path(sys.executable if getattr(sys, "frozen", False) else __file__).resolve().parent
URL_PLACES = "https://places.googleapis.com/v1/places:searchText"
CAMPOS_PLACES = ",".join([
    "places.id", "places.displayName", "places.formattedAddress",
    "places.nationalPhoneNumber", "places.internationalPhoneNumber",
    "places.websiteUri", "places.googleMapsUri", "places.rating",
    "places.userRatingCount", "places.businessStatus", "nextPageToken",
])

CONFIG_POR_DEFECTO = {
    "paginas_max": 3,
    "puntaje_minimo": 4,
    "remitente_nombre": "TU_NOMBRE",
    "remitente_descripcion": "desarrollador web freelance",
    "remitente_contacto": "TU_CORREO_O_WHATSAPP",
    "servicios": "diseño de páginas web y apps a medida",
}

# Enlaces que no cuentan como web propia (redes sociales y apps de terceros).
REDES = re.compile(
    r"(facebook\.com|fb\.com|instagram\.com|linktr\.ee|wa\.me|whatsapp\.com|tiktok\.com|"
    r"linkedin\.com|business\.site|sites\.google\.com|agendapro|weibook|schedulista|"
    r"tubarberia\.uy|guiadeo\.com)",
    re.I,
)

# País (en minúsculas y sin tildes) -> (código de región para Google, prefijos de celular).
# Sin prefijos no se puede saber si un número es celular (por ejemplo en México).
PAISES = {
    "uruguay": ("UY", ("5989",)),
    "argentina": ("AR", ("549",)),
    "chile": ("CL", ("569",)),
    "paraguay": ("PY", ("5959",)),
    "colombia": ("CO", ("573",)),
    "peru": ("PE", ("519",)),
    "espana": ("ES", ("346", "347")),
    "mexico": ("MX", ()),
    "estados unidos": ("US", ()),
}

PUNTOS_ESTADO = {
    "No tiene": 4,
    "Solo redes o app de terceros": 5,
    "Caída o no responde": 4,
    "Vieja": 3,
    "Mejorable": 1,
    "Correcta": 0,
}

COLUMNAS = [
    "Nombre", "Rubro", "Dirección", "Teléfono", "WhatsApp", "Correo", "Estado web",
    "Problema principal", "Web actual", "Problemas detectados", "Puntaje", "Prioridad",
    "Motivo", "Rating Google", "Reseñas", "Google Maps", "Plantilla asunto",
    "Plantilla correo", "Plantilla WhatsApp", "Fecha",
]


# ---------------------------------------------------------------- configuración

def cargar_config(ruta):
    if not ruta.exists():
        sys.exit(
            f"No encontré {ruta}.\n"
            "Copiá config.example.json como config.json y poné tu clave de Google Places."
        )
    with open(ruta, encoding="utf-8") as f:
        config = {**CONFIG_POR_DEFECTO, **json.load(f)}
    clave = str(config.get("google_api_key", "")).strip()
    if not clave or clave.startswith("TU_"):
        sys.exit(f"Falta la clave: poné tu clave de Google Places en google_api_key dentro de {ruta}.")
    return config


def datos_pais(pais):
    """Devuelve (código de región, prefijos de celular); (None, ()) si el país no está en PAISES."""
    clave = unicodedata.normalize("NFKD", pais).encode("ascii", "ignore").decode().strip().lower()
    return PAISES.get(clave, (None, ()))


def es_celular(telefono_internacional, prefijos):
    tel = re.sub(r"\D", "", telefono_internacional or "")
    return any(tel.startswith(p) for p in prefijos)


# ---------------------------------------------------------------- Google Places

def buscar_en_places(rubro, ciudad, pais, config, sesion):
    """Devuelve la lista de lugares crudos de Google Places (hasta paginas_max páginas)."""
    headers = {
        "Content-Type": "application/json",
        "X-Goog-Api-Key": config["google_api_key"],
        "X-Goog-FieldMask": CAMPOS_PLACES,
    }
    cuerpo = {
        "textQuery": f"{rubro} en {ciudad}, {pais}",
        "languageCode": "es",
        "pageSize": 20,
    }
    region, _ = datos_pais(pais)
    if region:
        cuerpo["regionCode"] = region
    lugares = []
    for pagina in range(int(config["paginas_max"])):
        try:
            r = sesion.post(URL_PLACES, headers=headers, json=cuerpo, timeout=30)
        except requests.RequestException as e:
            sys.exit(f"No pude conectarme a Google Places. Revisá tu conexión a internet.\n({e})")
        if r.status_code in (400, 401, 403):
            sys.exit(
                f"Google rechazó el pedido ({r.status_code}). Revisá que la clave en config.json esté bien copiada\n"
                f"y que tengas activada la Places API (New) en Google Cloud.\n\nDetalle: {r.text[:300]}"
            )
        if r.status_code != 200:
            sys.exit(f"Google Places respondió {r.status_code}: {r.text[:500]}")
        datos = r.json()
        lugares.extend(datos.get("places", []))
        print(f"  Página {pagina + 1}: {len(datos.get('places', []))} negocios")
        token = datos.get("nextPageToken")
        if not token:
            break
        cuerpo["pageToken"] = token
        time.sleep(2)
    return lugares


def separar_negocios(lugares, rubro, pais):
    """Un dict por negocio abierto, sin repetidos, con el tipo de web."""
    vistos = set()
    negocios = []
    for p in lugares:
        if p.get("id") in vistos:
            continue
        vistos.add(p.get("id"))
        if p.get("businessStatus") and p["businessStatus"] != "OPERATIONAL":
            continue
        web = p.get("websiteUri") or ""
        tipo_web = "sin_web"
        if web:
            tipo_web = "solo_redes_o_terceros" if REDES.search(web) else "web_propia"
        negocios.append({
            "place_id": p.get("id"),
            "nombre": (p.get("displayName") or {}).get("text", ""),
            "rubro": rubro,
            "pais": pais,
            "direccion": p.get("formattedAddress", ""),
            "telefono": p.get("nationalPhoneNumber", ""),
            "telefono_internacional": p.get("internationalPhoneNumber", ""),
            "web": web,
            "tipo_web": tipo_web,
            "google_maps": p.get("googleMapsUri", ""),
            "rating": p.get("rating"),
            "cantidad_resenas": p.get("userRatingCount", 0) or 0,
        })
    return negocios


# ---------------------------------------------------------------- revisión de webs

def descargar_web(url, sesion):
    """Devuelve (codigo, html, url_final). Si falla, codigo es 0."""
    try:
        r = sesion.get(
            url, timeout=15, allow_redirects=True,
            headers={"User-Agent": "Mozilla/5.0 (revisión de sitio web)"},
        )
        return r.status_code, r.text or "", r.url or url
    except requests.RequestException:
        return 0, "", url


def analizar_web(negocio, codigo, html, url_final, anio_actual=None):
    anio_actual = anio_actual or datetime.date.today().year
    problemas = []
    if not codigo or codigo >= 400 or len(html) < 200:
        estado = "Caída o no responde"
        problemas.append(f"El sitio no carga o da error ({codigo or 'sin respuesta'})")
    else:
        if not url_final.lower().startswith("https:") and not negocio["web"].lower().startswith("https:"):
            problemas.append('No usa HTTPS (el navegador lo marca como "no seguro")')
        if not re.search(r"<meta[^>]+name=[\"']viewport[\"']", html, re.I):
            problemas.append("No está adaptado a celulares (falta meta viewport)")
        anios = [int(m) for m in re.findall(r"(?:©|&copy;|copyright)\s*(?:\d{4}\s*[-–]\s*)?(\d{4})", html, re.I)]
        if anios and anio_actual - max(anios) >= 3:
            problemas.append(f"Pie de página con año {max(anios)} (parece abandonado)")
        if re.search(r"<(table[^>]+width=|font[\s>]|marquee|frameset)", html, re.I) or re.search(r"\.swf[\"']", html, re.I):
            problemas.append("Usa tecnología web antigua (tablas, <font>, Flash)")
        if re.search(r"jquery[-.]?1\.\d", html, re.I):
            problemas.append("Usa librerías muy viejas (jQuery 1.x)")
        if not re.search(r"wa\.me|api\.whatsapp\.com", html, re.I):
            problemas.append("No tiene botón de WhatsApp")
        if not re.search(r"<form", html, re.I):
            problemas.append("No tiene formulario de contacto o reservas")
        if not re.search(r"<meta[^>]+name=[\"']description[\"']", html, re.I):
            problemas.append("Sin descripción para Google (SEO básico)")
        graves = sum(1 for p in problemas if re.search(r"celulares|abandonado|antigua|viejas|HTTPS", p))
        estado = "Vieja" if graves >= 2 else "Mejorable" if len(problemas) >= 3 else "Correcta"

    correos = []
    for c in re.findall(r"[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}", html, re.I):
        c = c.lower()
        if re.search(r"\.(png|jpe?g|gif|webp|svg)$", c) or re.search(r"(sentry|wixpress|example|domain\.com)", c):
            continue
        if c not in correos:
            correos.append(c)
    return {**negocio, "estado_web": estado, "problemas_web": problemas, "correo": correos[0] if correos else ""}


def marcar_sin_web(negocio):
    solo_redes = negocio["tipo_web"] == "solo_redes_o_terceros"
    return {
        **negocio,
        "estado_web": "Solo redes o app de terceros" if solo_redes else "No tiene",
        "problemas_web": [f"Su único enlace es {negocio['web']} (no es una web propia)"] if solo_redes
        else ["No tiene página web en su ficha de Google Maps"],
        "correo": "",
    }


def revisar_webs(negocios, sesion):
    con_web = [n for n in negocios if n["tipo_web"] == "web_propia"]
    sin_web = [marcar_sin_web(n) for n in negocios if n["tipo_web"] != "web_propia"]
    print(f"  Revisando {len(con_web)} webs propias…")
    with ThreadPoolExecutor(max_workers=5) as pool:
        respuestas = list(pool.map(lambda n: descargar_web(n["web"], sesion), con_web))
    revisados = [analizar_web(n, *resp) for n, resp in zip(con_web, respuestas)]
    return revisados + sin_web


# ---------------------------------------------------------------- puntaje

def puntuar(n):
    p = PUNTOS_ESTADO.get(n["estado_web"], 0)
    motivo = [n["estado_web"]]
    if n["telefono"]:
        p += 1
        motivo.append("tiene teléfono")
    else:
        p -= 2
        motivo.append("sin teléfono")
    if es_celular(n.get("telefono_internacional"), datos_pais(n["pais"])[1]):
        p += 1
        motivo.append("celular (WhatsApp probable)")
    if n["cantidad_resenas"] >= 20:
        p += 1
        motivo.append(f"{n['cantidad_resenas']} reseñas (negocio activo)")
    if n["correo"]:
        p += 1
        motivo.append("tiene correo")
    prioridad = "Alta" if p >= 6 else "Media" if p >= 4 else "Baja" if p >= 2 else "Descartar"
    return {**n, "puntaje": p, "prioridad": prioridad, "motivo": ", ".join(motivo)}


def puntuar_y_filtrar(negocios, puntaje_minimo):
    puntuados = [puntuar(n) for n in negocios]
    elegidos = [n for n in puntuados if n["puntaje"] >= int(puntaje_minimo)]
    return sorted(elegidos, key=lambda n: n["puntaje"], reverse=True)


# ---------------------------------------------------------------- plantillas (editables)

def problema_principal(n):
    """Un problema principal por negocio, con la frase que se usa en los mensajes."""
    p = " ".join(n.get("problemas_web") or [])
    if n["estado_web"] == "No tiene":
        return "Sin web", "no tiene página web, así que cuando alguien lo busca en Google solo encuentra la ficha de Maps"
    if n["estado_web"] == "Solo redes o app de terceros":
        dominio = re.sub(r"^https?://(www\.)?", "", n["web"]).split("/")[0]
        return "Solo redes", f"no tiene web propia y depende de {dominio} para que lo encuentren"
    if n["estado_web"] == "Caída o no responde":
        return "Web caída", f"tiene la página web ({n['web']}) caída, y los clientes que llegan desde Google se van"
    if "celulares" in p:
        return "Sin celular", "tiene una página web que no se adapta a celulares, que es desde donde entra la mayoría de la gente"
    if "HTTPS" in p:
        return "Sin HTTPS", 'tiene una página web sin HTTPS, y el navegador la marca como "no segura"'
    if re.search(r"abandonado|antigua|viejas", p):
        return "Web vieja", "tiene una página web que se ve desactualizada y transmite que el negocio está menos activo de lo que está"
    faltas = " y ".join(re.sub(r"^No tiene ", "", x).lower() for x in (n.get("problemas_web") or [])[:2])
    return "Mejorable", f"tiene una página web a la que le faltan cosas que hoy ayudan a vender, como {faltas}"


def armar_fila(n, config, ciudad, hoy):
    firma = f"{config['remitente_nombre']} · {config['remitente_contacto']}"
    presentacion = f"Soy {config['remitente_nombre']}, {config['remitente_descripcion']} en {ciudad}"
    clave, frase = problema_principal(n)
    celular = es_celular(n.get("telefono_internacional"), datos_pais(n["pais"])[1])

    asunto = f"Una idea para la web de {n['nombre']}"
    correo = (
        f"Hola, ¿cómo está?\n\n"
        f"{presentacion}. Vi que {n['nombre']} {frase}.\n\n"
        f"Trabajo con negocios de rubros como {n['rubro']} y les armo páginas web y apps a medida: "
        f"rápidas, que se ven bien en el celular, con botón de WhatsApp y formulario de contacto o reservas.\n\n"
        f"Si le interesa, le preparo una propuesta sin compromiso. ¿Le parece si lo charlamos?\n\n"
        f"Saludos,\n{firma}"
    )
    whatsapp = (
        f"Hola, ¿qué tal? {presentacion}. Vi que {n['nombre']} {frase}. "
        f"Hago webs para rubros como {n['rubro']}, con WhatsApp directo y pensadas para el celular. "
        f"¿Te interesa que te muestre un ejemplo?"
    )
    return {
        "Nombre": n["nombre"],
        "Rubro": n["rubro"],
        "Dirección": n["direccion"],
        "Teléfono": n["telefono"],
        "WhatsApp": n["telefono"] if celular else "No figura",
        "Correo": n["correo"] or "No figura",
        "Estado web": n["estado_web"],
        "Problema principal": clave,
        "Web actual": n["web"],
        "Problemas detectados": "; ".join(n.get("problemas_web") or []),
        "Puntaje": n["puntaje"],
        "Prioridad": n["prioridad"],
        "Motivo": n["motivo"],
        "Rating Google": n["rating"],
        "Reseñas": n["cantidad_resenas"],
        "Google Maps": n["google_maps"],
        "Plantilla asunto": asunto,
        "Plantilla correo": correo,
        "Plantilla WhatsApp": whatsapp,
        "Fecha": hoy,
    }


# ---------------------------------------------------------------- Excel

def exportar_excel(filas, ruta):
    wb = Workbook()
    ws = wb.active
    ws.title = "Posibles clientes"
    ws.append(COLUMNAS)
    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="305496")
    for fila in filas:
        ws.append([fila[c] for c in COLUMNAS])
    anchos = {"Nombre": 30, "Dirección": 35, "Problemas detectados": 45, "Motivo": 40,
              "Google Maps": 30, "Plantilla asunto": 35, "Plantilla correo": 60, "Plantilla WhatsApp": 60}
    for i, col in enumerate(COLUMNAS, start=1):
        letra = ws.cell(row=1, column=i).column_letter
        ws.column_dimensions[letra].width = anchos.get(col, 16)
    col_maps = COLUMNAS.index("Google Maps") + 1
    for fila in ws.iter_rows(min_row=2):
        for celda in fila:
            celda.alignment = Alignment(wrap_text=True, vertical="top")
        maps = fila[col_maps - 1]
        if maps.value:
            maps.hyperlink = maps.value
            maps.style = "Hyperlink"
            maps.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(ruta)


# ---------------------------------------------------------------- programa

def ejecutar(rubro, ciudad, pais, config, sesion, carpeta_salida):
    hoy = datetime.date.today().isoformat()
    print(f"Buscando {rubro} en {ciudad}, {pais}…")
    lugares = buscar_en_places(rubro, ciudad, pais, config, sesion)
    negocios = separar_negocios(lugares, rubro, pais)
    print(f"  {len(negocios)} negocios abiertos")
    revisados = revisar_webs(negocios, sesion)
    elegidos = puntuar_y_filtrar(revisados, config["puntaje_minimo"])
    print(f"  {len(elegidos)} pasan el puntaje mínimo ({config['puntaje_minimo']})")
    filas = [armar_fila(n, config, ciudad, hoy) for n in elegidos]
    nombre = "clientes-" + re.sub(r"[^\w-]+", "-", f"{rubro}-{ciudad}-{pais}").strip("-").lower() + f"-{hoy}.xlsx"
    ruta = Path(carpeta_salida) / nombre
    exportar_excel(filas, ruta)
    print(f"Listo: {ruta}")
    return ruta


def preguntar(texto, por_defecto=""):
    extra = f" [{por_defecto}]" if por_defecto else ""
    return input(f"{texto}{extra}: ").strip() or por_defecto


def main():
    if len(sys.argv) == 1:
        modo_interactivo()
        return
    parser = argparse.ArgumentParser(description="Busca posibles clientes en Google Places y arma un Excel. No envía nada.")
    parser.add_argument("rubro", help='por ejemplo "carpinterías"')
    parser.add_argument("ciudad", help='por ejemplo "Córdoba"')
    parser.add_argument("pais", help='por ejemplo "Argentina"')
    parser.add_argument("--config", default=str(CARPETA / "config.json"), help="ruta al config.json")
    parser.add_argument("--salida", default=".", help="carpeta donde guardar el Excel")
    args = parser.parse_args()

    config = cargar_config(Path(args.config))
    ejecutar(args.rubro, args.ciudad, args.pais, config, requests.Session(), args.salida)


def preguntar_obligatorio(texto):
    respuesta = ""
    while not respuesta:
        respuesta = preguntar(texto)
    return respuesta


def modo_interactivo():
    """Para el doble clic: pregunta los datos y no cierra la ventana hasta apretar Enter."""
    try:
        config = cargar_config(CARPETA / "config.json")
        print("Buscar posibles clientes en Google Maps (no envía nada)\n")
        rubro = preguntar_obligatorio('Rubro, por ejemplo "carpinterías"')
        ciudad = preguntar_obligatorio("Ciudad")
        pais = preguntar_obligatorio("País")
        ruta = ejecutar(rubro, ciudad, pais, config, requests.Session(), CARPETA)
        print(f"\nAbrí el Excel: {ruta.name}")
    except SystemExit as e:
        if e.code not in (None, 0):
            print(f"\n{e.code}")
    except Exception as e:  # que la ventana no se cierre sin mostrar el error
        print(f"\nError inesperado: {e}")
    input("\nPresioná Enter para cerrar…")


if __name__ == "__main__":
    main()
